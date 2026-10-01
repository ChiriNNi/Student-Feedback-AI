import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import Text, cast, func, select
from sqlalchemy.orm import Session

from ..db import get_db
from ..models import KINDS, SENTIMENTS, Course, Feedback, Service, TopicCategory, User
from ..nlp.analyzer import label_for
from ..security import admin_only, staff_only
from ..services import (
    Filters, analyze_into, apply_filters, get_matchers, recompute_alerts, scope_query, semester_key,
)
from .analytics import filters_dep

router = APIRouter(prefix="/api/feedback", tags=["feedback"])


def fb_out(fb: Feedback) -> dict:
    return {
        "id": fb.id, "kind": fb.kind, "semester": fb.semester,
        "course": f"{fb.course.code} — {fb.course.title}" if fb.course else None,
        "course_id": fb.course_id,
        "service": fb.service.name if fb.service else None, "service_id": fb.service_id,
        "department": fb.department.name if fb.department else None,
        "rating": fb.rating, "text": fb.text, "language": fb.language,
        "sentiment": fb.sentiment, "sentiment_score": fb.sentiment_score, "confidence": fb.confidence,
        "topics": fb.topics or [], "aspects": fb.aspects or {}, "keywords": fb.keywords or [],
        "pii_masked": fb.pii_masked, "is_corrected": fb.is_corrected, "source": fb.source,
        "created_at": fb.created_at.isoformat() if fb.created_at else None,
    }


def _query(user: User, f: Filters):
    stmt = apply_filters(scope_query(select(Feedback), user), f)
    if f.topic:
        # topics is a JSON array; matching its text form keeps the query simple and portable
        stmt = stmt.where(cast(Feedback.topics, Text).contains(f'"{f.topic}"'))
    return stmt


@router.get("")
def list_feedback(
    page: int = Query(1, ge=1), size: int = Query(25, ge=1, le=200),
    sort: str = Query("newest", pattern="^(newest|oldest|most_negative|most_positive)$"),
    f: Filters = Depends(filters_dep), user: User = Depends(staff_only), db: Session = Depends(get_db),
):
    stmt = _query(user, f)
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    order = {
        "newest": Feedback.id.desc(), "oldest": Feedback.id.asc(),
        "most_negative": Feedback.sentiment_score.asc(), "most_positive": Feedback.sentiment_score.desc(),
    }[sort]
    items = db.scalars(stmt.order_by(order).offset((page - 1) * size).limit(size)).all()
    return {"total": total, "page": page, "size": size, "items": [fb_out(x) for x in items]}


class CorrectionIn(BaseModel):
    sentiment: str | None = None
    topics: list[str] | None = None


@router.patch("/{fb_id}", dependencies=[Depends(admin_only)])
def correct(fb_id: int, body: CorrectionIn, db: Session = Depends(get_db)):
    fb = db.get(Feedback, fb_id)
    if not fb:
        raise HTTPException(404, "Feedback not found")
    if body.sentiment:
        if body.sentiment not in SENTIMENTS:
            raise HTTPException(422, "Invalid sentiment")
        if label_for(fb.sentiment_score) != body.sentiment:
            fb.sentiment_score = {"positive": 0.6, "neutral": 0.0, "negative": -0.6}[body.sentiment]
        fb.sentiment = body.sentiment
    if body.topics is not None:
        known = set(db.scalars(select(TopicCategory.name)))
        topics = [t for t in body.topics if t in known]
        fb.topics = topics
        aspects = dict(fb.aspects or {})
        fb.aspects = {t: aspects.get(t, fb.sentiment) for t in topics}
    fb.is_corrected = True
    db.commit()
    return fb_out(fb)


@router.post("/reanalyze", dependencies=[Depends(admin_only)])
def reanalyze(db: Session = Depends(get_db)):
    """Re-run the NLP pipeline on all feedback that was not manually corrected (e.g. after editing topics)."""
    matchers = get_matchers(db)
    n = 0
    for fb in db.scalars(select(Feedback).where(Feedback.is_corrected.is_(False))).yield_per(500):
        analyze_into(fb, fb.text, matchers)
        n += 1
    db.commit()
    alerts = recompute_alerts(db)
    return {"reanalyzed": n, "alerts_created": alerts}


# ------------------------------------------------------------------ import

TEMPLATE_HEADERS = ["text", "rating", "semester", "kind", "course_code", "service"]


@router.get("/import/template", dependencies=[Depends(admin_only)])
def import_template():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADERS)
    w.writerow(["The lectures were clear but the deadlines overlap.", 4, "Fall 2026", "course", "CSS 101", ""])
    w.writerow(["Библиотека работает слишком мало по выходным.", 3, "Fall 2026", "service", "", "Library"])
    return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=feedback_import_template.csv"})


def _read_rows(name: str, data: bytes) -> list[dict]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it, [])]
        return [dict(zip(header, r)) for r in it]
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]


@router.post("/import", dependencies=[Depends(admin_only)])
async def import_feedback(file: UploadFile = File(...), default_semester: str = Query("Fall 2026"),
                          db: Session = Depends(get_db)):
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(413, "File too large (max 15 MB)")
    try:
        rows = _read_rows(file.filename or "upload.csv", data)
    except Exception:
        raise HTTPException(400, "Could not read the file. Use CSV (UTF-8) or XLSX with a header row.")
    if not rows:
        raise HTTPException(400, "The file has no data rows")
    if "text" not in rows[0]:
        raise HTTPException(400, "Missing required column 'text'. Download the template for the expected format.")

    courses = {c.code.lower(): c for c in db.scalars(select(Course))}
    services = {s.name.lower(): s for s in db.scalars(select(Service))}
    matchers = get_matchers(db)
    imported, errors = 0, []

    for i, r in enumerate(rows, start=2):
        text = str(r.get("text") or "").strip()
        if len(text) < 3:
            errors.append({"row": i, "error": "empty text"})
            continue
        rating = r.get("rating")
        try:
            rating = int(float(rating)) if rating not in (None, "") else None
            if rating is not None and not 1 <= rating <= 5:
                raise ValueError
        except (TypeError, ValueError):
            errors.append({"row": i, "error": f"rating must be 1-5 (got {r.get('rating')!r})"})
            continue
        semester = str(r.get("semester") or default_semester).strip()
        if semester_key(semester) == (0, 0):
            errors.append({"row": i, "error": f"bad semester {semester!r}"})
            continue
        code = str(r.get("course_code") or "").strip().lower()
        svc = str(r.get("service") or "").strip().lower()
        kind = str(r.get("kind") or "").strip().lower() or ("course" if code else "service" if svc else "survey")
        if kind not in KINDS:
            errors.append({"row": i, "error": f"kind must be one of {', '.join(KINDS)}"})
            continue
        fb = Feedback(kind=kind, semester=semester, rating=rating, source="import")
        if code:
            c = courses.get(code)
            if not c:
                errors.append({"row": i, "error": f"unknown course_code {r.get('course_code')!r}"})
                continue
            fb.course_id, fb.department_id = c.id, c.department_id
        if svc:
            s = services.get(svc)
            if not s:
                errors.append({"row": i, "error": f"unknown service {r.get('service')!r}"})
                continue
            fb.service_id = s.id
        analyze_into(fb, text, matchers)
        db.add(fb)
        imported += 1
    db.commit()
    alerts = recompute_alerts(db) if imported else 0
    return {"imported": imported, "skipped": len(errors), "errors": errors[:50], "alerts_created": alerts}


# ------------------------------------------------------------------ export

EXPORT_COLS = ["id", "semester", "kind", "department", "course", "service", "rating", "language",
               "sentiment", "sentiment_score", "topics", "text"]


def _export_rows(db, user, f):
    items = db.scalars(_query(user, f).order_by(Feedback.id)).all()
    for fb in items:
        d = fb_out(fb)
        d["topics"] = ", ".join(d["topics"])
        yield [d.get(c) for c in EXPORT_COLS]


@router.get("/export.csv")
def export_csv(f: Filters = Depends(filters_dep), user: User = Depends(staff_only), db: Session = Depends(get_db)):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPORT_COLS)
    w.writerows(_export_rows(db, user, f))
    stamp = datetime.now().strftime("%Y%m%d")
    return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=feedback_{stamp}.csv"})


@router.get("/export.xlsx")
def export_xlsx(f: Filters = Depends(filters_dep), user: User = Depends(staff_only), db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"
    ws.append([c.replace("_", " ").title() for c in EXPORT_COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", fgColor="4B3AA8")
    fills = {"positive": "D9F7E8", "negative": "FDE0E6", "neutral": "FFF4D6"}
    for row in _export_rows(db, user, f):
        ws.append(row)
        s = row[EXPORT_COLS.index("sentiment")]
        ws.cell(ws.max_row, EXPORT_COLS.index("sentiment") + 1).fill = PatternFill("solid", fgColor=fills.get(s, "FFFFFF"))
    widths = [8, 13, 10, 26, 34, 20, 8, 9, 11, 10, 34, 90]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=feedback_{stamp}.xlsx"})
